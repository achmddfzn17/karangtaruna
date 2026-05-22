"""Generate KEBUTUHAN_FUNGSIONAL.docx dari konten dokumen."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Calibri'
    return h


def add_para(doc, text, bold=False, italic=False, size=11, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        run = p.runs[0] if p.runs else p.add_run('')
        run.text = item
        run.font.name = 'Calibri'
        run.font.size = Pt(11)


def add_kv_table(doc, rows):
    """Tabel 2 kolom (Field | Value) untuk format fitur."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for i, (key, val) in enumerate(rows):
        cell_k = table.cell(i, 0)
        cell_v = table.cell(i, 1)
        cell_k.width = Cm(4.5)
        cell_v.width = Cm(11.5)

        cell_k.text = ''
        p_k = cell_k.paragraphs[0]
        run_k = p_k.add_run(key)
        run_k.bold = True
        run_k.font.name = 'Calibri'
        run_k.font.size = Pt(10)
        set_cell_shading(cell_k, 'E7E6E6')
        set_cell_borders(cell_k)

        cell_v.text = ''
        p_v = cell_v.paragraphs[0]
        # Jika value berupa list (untuk alur proses), buat numbered list
        if isinstance(val, list):
            for idx, step in enumerate(val):
                if idx == 0:
                    run_v = p_v.add_run(f"{idx + 1}. {step}")
                else:
                    new_p = cell_v.add_paragraph()
                    run_v = new_p.add_run(f"{idx + 1}. {step}")
                run_v.font.name = 'Calibri'
                run_v.font.size = Pt(10)
        else:
            run_v = p_v.add_run(val)
            run_v.font.name = 'Calibri'
            run_v.font.size = Pt(10)
        set_cell_borders(cell_v)
    return table


def add_actor_table(doc):
    headers = ['Fitur / Halaman', 'Pengunjung', 'Anggota', 'Admin', 'Super Admin']
    rows = [
        ['Halaman Publik (berita, artikel, galeri, program)', 'Ya', 'Ya', 'Ya', 'Ya'],
        ['SUS Survey', 'Ya', 'Ya', 'Ya', 'Ya'],
        ['Aspirasi (kirim)', 'Ya', 'Ya', 'Ya', 'Ya'],
        ['Login / Registrasi', 'Ya', '-', '-', '-'],
        ['Member Dashboard (profil, kegiatan, sertifikat)', 'Tidak', 'Ya', 'Ya', 'Ya'],
        ['E-Voting (memberikan suara)', 'Tidak', 'Ya', 'Tidak', 'Tidak'],
        ['Lihat Iuran & Keuangan Pribadi', 'Tidak', 'Ya', 'Ya', 'Ya'],
        ['Dashboard Admin (backoffice)', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['Manajemen Anggota', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['Manajemen Kegiatan & Absensi', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['Penerbitan Sertifikat', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['CMS (Berita, Artikel, Galeri)', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['Manajemen Iuran & Keuangan', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['Kelola Aspirasi & Voting', 'Tidak', 'Tidak', 'Ya', 'Ya'],
        ['Kelola Admin & Audit Log', 'Tidak', 'Tidak', 'Tidak', 'Ya'],
    ]

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    # Header
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, '4472C4')
        for r in cell.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_borders(cell)
    # Body
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            if j > 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_borders(cell)


def add_summary_table(doc, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    headers = ['Kode', 'Nama Fitur', 'Aktor Utama']
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, '4472C4')
        for r in cell.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_borders(cell)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            if j == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_borders(cell)


# =========================================================
# DATA FITUR
# =========================================================

FITUR = [
    {
        'kode': 'F001',
        'nama': 'Fitur Autentikasi (Login & Logout)',
        'aktor': 'Anggota / Admin / Super Admin',
        'deskripsi': 'Sistem menyediakan autentikasi berbasis NextAuth dengan email & password. Role pengguna menentukan halaman tujuan setelah login.',
        'alur': [
            'Pengguna membuka halaman login.',
            'Memasukkan email dan password.',
            'Sistem memvalidasi credentials & status akun.',
            'Sistem membuat session berisi user_id dan role.',
            'Anggota diarahkan ke /member/dashboard, Admin/Super Admin ke /dashboard.',
        ],
        'kriteria': [
            'Login berhasil hanya dengan credentials valid.',
            'Pesan error informatif untuk input salah.',
            'Redirect sesuai role.',
            'Session aktif hingga logout / expired.',
        ],
    },
    {
        'kode': 'F002',
        'nama': 'Fitur Registrasi Anggota',
        'aktor': 'Pengunjung (calon anggota)',
        'deskripsi': 'Calon anggota dapat mendaftarkan diri secara mandiri. Akun baru dibuat dengan role ANGGOTA dan terhubung dengan data biodata anggota.',
        'alur': [
            'Pengunjung membuka halaman registrasi.',
            'Mengisi form: nama, email, password, NIK, tempat/tanggal lahir, jenis kelamin, alamat, no. HP.',
            'Sistem memvalidasi keunikan email & NIK.',
            'Sistem membuat record User (role ANGGOTA) dan Anggota dengan status AKTIF.',
            'Anggota dapat langsung login.',
        ],
        'kriteria': [
            'Email & NIK unik.',
            'Password ter-hash menggunakan bcrypt.',
            'Validasi format input berjalan.',
            'Data tersimpan di tabel users dan anggota.',
        ],
    },
    {
        'kode': 'F003',
        'nama': 'Fitur Manajemen Anggota (Backoffice)',
        'aktor': 'Admin & Super Admin',
        'deskripsi': 'Admin dapat melakukan CRUD data anggota, mengubah status (Aktif / Non-Aktif / Alumni), serta membuatkan akun login untuk anggota yang belum memiliki akun.',
        'alur': [
            'Admin membuka menu Anggota.',
            'Sistem menampilkan daftar anggota dengan filter status.',
            'Admin dapat menambah, mengedit, atau menghapus anggota.',
            'Admin dapat mengakses halaman "Buat Akun" untuk anggota yang belum punya userId.',
            'Sistem menyimpan perubahan ke tabel anggota.',
        ],
        'kriteria': [
            'NIK & email anggota unik.',
            'Status anggota dapat diubah.',
            'Foto anggota tersimpan di Supabase Storage.',
            'Akun login dapat dibuatkan oleh admin untuk anggota existing.',
        ],
    },
    {
        'kode': 'F004',
        'nama': 'Fitur Manajemen Kegiatan',
        'aktor': 'Admin & Super Admin',
        'deskripsi': 'Admin dapat membuat, mengedit, dan menghapus kegiatan. Setiap kegiatan memiliki jenis (Sosial, Pendidikan, Ekonomi, Olahraga, Seni Budaya, Lainnya), status (Upcoming / Ongoing / Selesai / Dibatalkan), tanggal, lokasi, anggaran, dan thumbnail.',
        'alur': [
            'Admin membuka menu Kegiatan.',
            'Mengisi form: nama, deskripsi, jenis, tanggal mulai/selesai, lokasi, anggaran, thumbnail.',
            'Admin menambahkan peserta dari daftar anggota.',
            'Status kegiatan dapat di-update sesuai progres.',
        ],
        'kriteria': [
            'Kegiatan tersimpan di tabel kegiatan.',
            'Peserta tersimpan di anggota_kegiatan (unik per anggota-kegiatan).',
            'Thumbnail tersimpan di storage.',
            'Status kegiatan dapat diperbarui.',
        ],
    },
    {
        'kode': 'F005',
        'nama': 'Fitur Absensi Kegiatan',
        'aktor': 'Admin & Super Admin',
        'deskripsi': 'Sistem mencatat kehadiran anggota pada setiap kegiatan. Status hadir pada peserta menjadi syarat penerbitan sertifikat.',
        'alur': [
            'Admin membuka halaman peserta kegiatan.',
            'Sistem menampilkan daftar anggota terdaftar.',
            'Admin mencentang anggota yang hadir.',
            'Sistem mengupdate field hadir di anggota_kegiatan.',
        ],
        'kriteria': [
            'Hanya peserta terdaftar yang dapat ditandai hadir.',
            'Perubahan tersimpan real-time.',
            'Data kehadiran menjadi sumber data untuk laporan & sertifikat.',
        ],
    },
    {
        'kode': 'F006',
        'nama': 'Fitur Penerbitan Sertifikat dengan QR Code',
        'aktor': 'Admin & Super Admin (penerbitan); Publik (verifikasi)',
        'deskripsi': 'Sistem menerbitkan sertifikat keikutsertaan otomatis untuk anggota yang hadir di sebuah kegiatan, lengkap dengan nomor sertifikat unik dan QR code yang mengarah ke halaman verifikasi /verify/[nomor]. Sertifikat dapat diunduh sebagai PDF.',
        'alur': [
            'Admin membuka menu Sertifikat.',
            'Memilih kegiatan dengan peserta yang sudah hadir.',
            'Sistem generate nomor sertifikat unik & QR code.',
            'Sistem menyimpan record di tabel sertifikat.',
            'Sertifikat dapat di-generate ke PDF (jsPDF) untuk diunduh atau dibagikan ke anggota.',
            'Pihak ketiga dapat memindai QR code untuk verifikasi keaslian.',
        ],
        'kriteria': [
            'Nomor sertifikat unik per kegiatan-anggota.',
            'QR code terbaca dan mengarah ke halaman verifikasi.',
            'Halaman verifikasi menampilkan nama anggota, kegiatan, dan tanggal.',
            'PDF dapat diunduh.',
        ],
    },
    {
        'kode': 'F007',
        'nama': 'Fitur Manajemen Iuran Anggota',
        'aktor': 'Admin & Super Admin (input); Anggota (lihat)',
        'deskripsi': 'Sistem mencatat pembayaran iuran bulanan setiap anggota. Anggota dapat melihat status iuran pribadinya, sedangkan admin melihat rekap iuran seluruh anggota per bulan/tahun.',
        'alur': [
            'Admin membuka menu Iuran.',
            'Memilih anggota, bulan, tahun, dan jumlah pembayaran.',
            'Sistem menyimpan ke tabel iuran_anggota (unik per anggota-bulan-tahun).',
            'Anggota dapat melihat riwayat iuran pribadi dan tunggakan di member dashboard.',
        ],
        'kriteria': [
            'Tidak boleh ada duplikasi iuran (anggota + bulan + tahun harus unik).',
            'Rekap iuran dapat difilter per bulan/tahun.',
            'Anggota hanya melihat iuran miliknya.',
        ],
    },
    {
        'kode': 'F008',
        'nama': 'Fitur Manajemen Keuangan Organisasi',
        'aktor': 'Admin & Super Admin',
        'deskripsi': 'Sistem mencatat transaksi keuangan organisasi (kas masuk dan keluar) dengan kategorisasi. Sistem menampilkan ringkasan saldo, total pemasukan, dan total pengeluaran beserta grafik tren bulanan.',
        'alur': [
            'Admin membuka menu Keuangan.',
            'Mengisi form transaksi: tanggal, jenis (MASUK/KELUAR), kategori, jumlah, keterangan, bukti (opsional).',
            'Sistem menyimpan ke tabel transaksi_keuangan.',
            'Dashboard menampilkan ringkasan & grafik (Recharts).',
            'Laporan dapat diekspor ke Excel/PDF.',
        ],
        'kriteria': [
            'Transaksi terklasifikasi MASUK/KELUAR.',
            'Saldo dihitung otomatis.',
            'Bukti transaksi tersimpan di storage.',
            'Ekspor laporan berjalan.',
        ],
    },
    {
        'kode': 'F009',
        'nama': 'Fitur CMS Berita & Artikel',
        'aktor': 'Admin & Super Admin (kelola); Publik (baca)',
        'deskripsi': 'Sistem menyediakan CMS berita & artikel dengan editor WYSIWYG (TipTap). Konten memiliki status (DRAFT / PUBLISHED / ARCHIVED), slug unik, tags, kategori, dan view counter. Hanya konten PUBLISHED yang muncul di halaman publik.',
        'alur': [
            'Admin membuka menu Berita atau Artikel.',
            'Mengisi form: judul, isi (TipTap), ringkasan, thumbnail, tags, kategori, status.',
            'Sistem auto-generate slug dari judul.',
            'Konten dengan status PUBLISHED otomatis tampil di /berita dan /artikel.',
            'View count bertambah otomatis saat dikunjungi.',
        ],
        'kriteria': [
            'Slug unik.',
            'Konten DRAFT/ARCHIVED tidak tampil di publik.',
            'SEO-friendly URL.',
            'Sanitasi HTML dengan DOMPurify.',
            'Gambar di-upload ke Supabase Storage.',
        ],
    },
    {
        'kode': 'F010',
        'nama': 'Fitur Galeri Kegiatan',
        'aktor': 'Admin & Super Admin (upload); Publik & Anggota (lihat)',
        'deskripsi': 'Sistem menyediakan galeri foto dan video yang dapat dikaitkan dengan kegiatan tertentu. Pengunjung dapat melihat dokumentasi kegiatan organisasi.',
        'alur': [
            'Admin membuka menu Galeri.',
            'Upload file (foto/video) beserta judul, deskripsi, dan kegiatan terkait.',
            'File tersimpan di Supabase Storage.',
            'Galeri tampil di halaman publik /galeri dengan filter berdasarkan kegiatan.',
        ],
        'kriteria': [
            'Tipe item dapat berupa FOTO atau VIDEO.',
            'Dapat di-link ke kegiatanId.',
            'File tersimpan di storage.',
            'Galeri publik menampilkan grid responsif.',
        ],
    },
    {
        'kode': 'F011',
        'nama': 'Fitur Aspirasi (Kotak Saran)',
        'aktor': 'Pengunjung & Anggota (kirim); Admin (kelola)',
        'deskripsi': 'Sistem menyediakan kanal aspirasi/saran/keluhan untuk masyarakat dan anggota. Aspirasi memiliki status alur (PENDING -> DIPROSES -> SELESAI / DITOLAK) dan dapat dibalas oleh admin.',
        'alur': [
            'Pengguna mengisi form aspirasi: judul, pesan, kategori (nama opsional jika anonim).',
            'Sistem menyimpan ke tabel aspirasi dengan status PENDING.',
            'Admin meninjau di backoffice dan dapat membalas + mengubah status.',
            'Anggota terdaftar dapat melihat riwayat aspirasi pribadinya.',
        ],
        'kriteria': [
            'Aspirasi anonim diperbolehkan.',
            'Status alur berjalan.',
            'Admin dapat memberi balasan.',
            'Notifikasi terkirim ke anggota saat aspirasi dibalas.',
        ],
    },
    {
        'kode': 'F012',
        'nama': 'Fitur E-Voting (Polling)',
        'aktor': 'Super Admin / Admin (buat polling); Anggota (vote)',
        'deskripsi': 'Sistem menyediakan e-voting internal untuk pemilihan pengurus, keputusan organisasi, atau survei opini. Setiap anggota hanya dapat memilih satu opsi per polling. Hasil ditampilkan secara real-time.',
        'alur': [
            'Admin membuat polling: judul, deskripsi, daftar opsi, tanggal kadaluwarsa.',
            'Polling aktif tampil di member dashboard.',
            'Anggota memilih opsi & submit.',
            'Sistem mencegah double-vote (unique userId + pollingId).',
            'Hasil divisualisasikan dalam grafik.',
        ],
        'kriteria': [
            'Satu anggota satu suara per polling.',
            'Polling dengan expiresAt lewat tidak dapat di-vote.',
            'Hasil real-time.',
            'Hanya polling isActive yang tampil.',
        ],
    },
    {
        'kode': 'F013',
        'nama': 'Fitur Notifikasi',
        'aktor': 'Sistem (kirim); Anggota & Admin (terima)',
        'deskripsi': 'Sistem mengirimkan notifikasi internal kepada pengguna untuk peristiwa penting (kegiatan baru, sertifikat terbit, balasan aspirasi, polling baru, pengingat iuran).',
        'alur': [
            'Sistem trigger notifikasi saat event tertentu terjadi.',
            'Record disimpan di tabel notifications dengan field userId, title, message, type, isRead.',
            'User membuka menu Notifikasi & menandai sebagai dibaca.',
        ],
        'kriteria': [
            'Notifikasi tampil real-time di header (badge counter).',
            'Dapat ditandai sudah dibaca.',
            'Notifikasi broadcast (userId = null) tampil ke semua user.',
        ],
    },
    {
        'kode': 'F014',
        'nama': 'Fitur Survei SUS (System Usability Scale)',
        'aktor': 'Pengunjung & Anggota',
        'deskripsi': 'Sistem menyediakan survei SUS 10 pertanyaan untuk mengukur kebergunaan sistem. Skor SUS dihitung otomatis (0-100) dan dikategorikan (Excellent / Good / OK / Poor). Hasil agregat menjadi indikator UX.',
        'alur': [
            'Responden membuka halaman /sus.',
            'Mengisi 10 pertanyaan skala 1-5.',
            'Sistem menghitung skor: ((SUM ganjil - 5) + (25 - SUM genap)) x 2.5.',
            'Sistem menentukan kategori berdasarkan skor.',
            'Admin melihat rekap di backoffice (rata-rata skor, distribusi kategori).',
        ],
        'kriteria': [
            'Perhitungan skor SUS sesuai standar.',
            'Kategori otomatis.',
            'Admin dapat melihat agregat & detail per responden.',
            'Ekspor data tersedia.',
        ],
    },
    {
        'kode': 'F015',
        'nama': 'Fitur Kalender Kegiatan',
        'aktor': 'Anggota & Admin',
        'deskripsi': 'Sistem menampilkan seluruh kegiatan organisasi dalam tampilan kalender bulanan/mingguan. Anggota dapat melihat jadwal kegiatan mendatang dan mengkliknya untuk detail.',
        'alur': [
            'Pengguna membuka menu Kalender.',
            'Sistem menampilkan kegiatan sesuai tanggalMulai & tanggalSelesai.',
            'Klik event menampilkan detail kegiatan.',
            'Filter berdasarkan jenis kegiatan tersedia.',
        ],
        'kriteria': [
            'Tampilan kalender responsif.',
            'Event ditampilkan sesuai jenis (warna kategori).',
            'Navigasi bulan/minggu berjalan.',
            'Klik event mengarah ke detail.',
        ],
    },
    {
        'kode': 'F016',
        'nama': 'Fitur Member Dashboard',
        'aktor': 'Anggota',
        'deskripsi': 'Halaman utama anggota yang menampilkan ringkasan: profil, jumlah kegiatan diikuti, jumlah sertifikat, status iuran, polling aktif, notifikasi belum dibaca.',
        'alur': [
            'Anggota login & diarahkan ke /member/dashboard.',
            'Sistem mengambil data dari berbagai modul (anggota, kegiatan, sertifikat, iuran, polling, notifikasi).',
            'Anggota dapat navigasi ke menu lain dari dashboard.',
        ],
        'kriteria': [
            'Data real-time dari database.',
            'Widget hanya menampilkan data milik anggota login.',
            'Tampilan responsif.',
        ],
    },
    {
        'kode': 'F017',
        'nama': 'Fitur Dashboard Admin (Backoffice)',
        'aktor': 'Admin & Super Admin',
        'deskripsi': 'Halaman utama backoffice yang menampilkan statistik organisasi: total anggota aktif, kegiatan berlangsung, saldo kas, iuran bulan berjalan, aspirasi pending, dan grafik tren keuangan & keanggotaan.',
        'alur': [
            'Admin login & diarahkan ke /dashboard.',
            'Sistem mengagregasi data dari seluruh modul.',
            'Dashboard menampilkan kartu statistik & grafik (Recharts).',
            'Quick action menuju modul-modul utama.',
        ],
        'kriteria': [
            'Statistik real-time.',
            'Grafik responsif.',
            'Data aspirasi pending hanya tampil bagi yang berwenang.',
            'Tampilan menyesuaikan role.',
        ],
    },
    {
        'kode': 'F018',
        'nama': 'Fitur Manajemen Admin (Super Admin Only)',
        'aktor': 'Super Admin Only',
        'deskripsi': 'Fitur eksklusif Super Admin untuk mengelola akun admin: menambah admin baru, mengubah jabatan/NIP, dan menonaktifkan admin. Admin baru dapat langsung login dengan role ADMIN.',
        'alur': [
            'Super Admin membuka menu Kelola Admin.',
            'Menambah admin: input nama, email, password, NIP, jabatan, no. HP.',
            'Sistem membuat record User (role ADMIN) dan Admin.',
            'Super Admin dapat mengedit atau menonaktifkan admin existing.',
        ],
        'kriteria': [
            'Hanya Super Admin yang dapat akses (middleware).',
            'Email & NIP admin unik.',
            'Password ter-hash.',
            'Admin yang dinonaktifkan tidak dapat login.',
        ],
    },
    {
        'kode': 'F019',
        'nama': 'Fitur Audit Log (Super Admin Only)',
        'aktor': 'Super Admin',
        'deskripsi': 'Sistem mencatat seluruh aktivitas penting (CREATE, UPDATE, DELETE, LOGIN) yang dilakukan oleh admin & super admin sebagai jejak audit. Setiap log memuat userId, userName, action, module, targetId, detail, dan ipAddress.',
        'alur': [
            'Sistem otomatis mencatat log saat aksi penting terjadi.',
            'Super Admin membuka menu Log.',
            'Sistem menampilkan log dengan filter berdasarkan modul, user, dan rentang tanggal.',
            'Detail log dapat ditinjau untuk audit kepatuhan.',
        ],
        'kriteria': [
            'Log tersimpan di tabel audit_logs.',
            'Filter & pencarian berjalan.',
            'Log tidak dapat dihapus oleh siapapun (immutable).',
            'Ekspor log tersedia.',
        ],
    },
    {
        'kode': 'F020',
        'nama': 'Fitur Profil Anggota',
        'aktor': 'Anggota',
        'deskripsi': 'Anggota dapat melihat dan mengupdate data profil pribadinya, termasuk foto, alamat, no. HP, pekerjaan, dan pendidikan.',
        'alur': [
            'Anggota membuka menu Profil.',
            'Sistem menampilkan data dari tabel anggota & users.',
            'Anggota mengubah data yang diperbolehkan & menyimpan.',
            'Foto baru di-upload ke Supabase Storage.',
        ],
        'kriteria': [
            'NIK & email tidak dapat diubah oleh anggota.',
            'Foto profil tersimpan.',
            'Perubahan tercatat di updatedAt.',
        ],
    },
    {
        'kode': 'F021',
        'nama': 'Fitur Halaman Publik (Landing & Profil Organisasi)',
        'aktor': 'Pengunjung',
        'deskripsi': 'Halaman publik yang dapat diakses tanpa login: landing page dengan profil organisasi, halaman tentang, daftar program kerja, berita & artikel terbaru, galeri kegiatan, dan tautan ke fitur publik lainnya (aspirasi, SUS).',
        'alur': [
            'Pengunjung membuka domain utama.',
            'Landing page menampilkan hero, ringkasan organisasi, program unggulan, berita terbaru, dan galeri.',
            'Pengunjung dapat navigasi ke halaman tentang, program, berita, artikel, galeri, kegiatan, aspirasi, dan SUS.',
        ],
        'kriteria': [
            'Halaman publik dapat diakses tanpa autentikasi.',
            'SEO-friendly.',
            'Responsif (mobile-first).',
            'Konten dimuat dari database (CMS).',
        ],
    },
]


# =========================================================
# BUILD DOCUMENT
# =========================================================

doc = Document()

# Set default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# Margin
section = doc.sections[0]
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)

# ============= COVER / TITLE =============
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('KEBUTUHAN FUNGSIONAL SISTEM')
run.bold = True
run.font.size = Pt(18)
run.font.name = 'Calibri'

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Sistem Informasi Manajemen Karang Taruna Berbasis Web')
run.bold = True
run.font.size = Pt(14)
run.font.name = 'Calibri'

doc.add_paragraph()

# ============= LATAR BELAKANG =============
add_heading(doc, 'LATAR BELAKANG', level=1)

add_para(doc, (
    'Karang Taruna merupakan organisasi kepemudaan di tingkat desa/kelurahan yang berperan dalam '
    'pemberdayaan generasi muda melalui kegiatan sosial, pendidikan, ekonomi, dan seni budaya. '
    'Dalam menjalankan operasionalnya, organisasi ini mengelola berbagai aspek seperti keanggotaan, '
    'kegiatan, keuangan, iuran anggota, dokumentasi, hingga komunikasi internal.'
))

add_para(doc, (
    'Saat ini sebagian besar pengelolaan masih dilakukan secara manual menggunakan dokumen fisik, '
    'spreadsheet terpisah, dan grup pesan singkat. Akibatnya muncul beberapa permasalahan: data anggota '
    'sulit dilacak dan rawan duplikasi, dokumentasi kegiatan tersebar di berbagai media, pencatatan iuran '
    'dan keuangan tidak transparan, sertifikat keikutsertaan kegiatan harus dibuat manual satu per satu, '
    'serta tidak ada kanal resmi bagi anggota dan masyarakat untuk menyampaikan aspirasi maupun mengikuti '
    'pemilihan internal.'
))

add_para(doc, (
    'Berdasarkan permasalahan tersebut, dibutuhkan sebuah Sistem Informasi Manajemen Karang Taruna berbasis '
    'web yang dapat mengintegrasikan seluruh proses tersebut dalam satu platform terpadu. Sistem ini '
    'mengotomatisasi penerbitan sertifikat dengan QR code, menyediakan kanal aspirasi dan e-voting, '
    'mempublikasikan berita & kegiatan ke masyarakat luas, serta menyediakan laporan keuangan dan iuran '
    'yang transparan.'
))

add_para(doc, 'Permasalahan yang diselesaikan sistem ini:', bold=True)
add_bullets(doc, [
    'Pendataan anggota masih manual dan tersebar, sulit untuk pelacakan riwayat keaktifan',
    'Dokumentasi kegiatan dan publikasi informasi tidak terpusat',
    'Pencatatan iuran dan keuangan organisasi tidak transparan kepada anggota',
    'Penerbitan sertifikat keikutsertaan kegiatan harus dibuat manual',
    'Tidak ada kanal resmi untuk aspirasi anggota dan masyarakat',
    'Pemilihan internal organisasi (e-voting) belum terdigitalisasi',
    'Pengurus tidak memiliki dashboard terpusat untuk memantau aktivitas organisasi',
])

# ============= DESKRIPSI SISTEM =============
add_heading(doc, 'DESKRIPSI SISTEM', level=1)

add_para(doc, (
    'Sistem Informasi Karang Taruna adalah aplikasi berbasis web yang dirancang untuk mengelola seluruh '
    'operasional organisasi karang taruna secara terpadu. Sistem terdiri dari tiga area utama: halaman '
    'publik (informasi dan publikasi untuk masyarakat), area anggota (member dashboard untuk anggota '
    'terdaftar), dan dashboard admin (backoffice untuk pengurus). Fitur unggulan sistem mencakup penerbitan '
    'sertifikat otomatis dengan QR code verifikasi, e-voting internal, kanal aspirasi, manajemen keuangan & '
    'iuran, serta survei kebergunaan sistem (SUS).'
))

add_heading(doc, 'AKTOR & HAK AKSES', level=2)
add_actor_table(doc)

# ============= KEBUTUHAN FUNGSIONAL =============
doc.add_page_break()
add_heading(doc, 'KEBUTUHAN FUNGSIONAL', level=1)
add_para(doc, 'Berikut adalah daftar kebutuhan fungsional lengkap sistem Informasi Manajemen Karang Taruna:')

for f in FITUR:
    # Heading per fitur
    h = doc.add_paragraph()
    run = h.add_run(f"{f['kode']}  —  {f['nama']}")
    run.bold = True
    run.font.size = Pt(13)
    run.font.name = 'Calibri'

    rows = [
        ('Kode & Nama', f"{f['kode']} — {f['nama']}"),
        ('Aktor / Role', f['aktor']),
        ('Deskripsi', f['deskripsi']),
        ('Alur Proses', f['alur']),
        ('Kriteria Penerimaan', f['kriteria']),
    ]
    add_kv_table(doc, rows)
    doc.add_paragraph()

# ============= RINGKASAN =============
doc.add_page_break()
add_heading(doc, 'RINGKASAN KEBUTUHAN FUNGSIONAL', level=1)

summary_rows = [[f['kode'], f['nama'], f['aktor']] for f in FITUR]
add_summary_table(doc, summary_rows)

# Save
output_path = '/projects/sandbox/karangtaruna/KEBUTUHAN_FUNGSIONAL.docx'
doc.save(output_path)
print(f'OK: {output_path}')
